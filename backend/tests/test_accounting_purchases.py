"""Purchases in the accountant package: rows, files, grouping under card invoices, warnings."""

import zipfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from support import ACCOUNTANT, UserFactory
from time_reporting.core.config import get_settings
from time_reporting.core.cqrs import Bus
from time_reporting.modules.accounting.contracts import (
    BuildAccountantPackage,
    GetAccountantPackageStatus,
)
from time_reporting.modules.purchases.contracts import (
    AddPurchaseDocument,
    CardReceiptLink,
    LinkCardReceipts,
    PaymentMethod,
    PaymentStatus,
    PurchaseDetails,
    PurchaseDocumentDTO,
    PurchaseKind,
    RegisterPurchaseDocument,
)
from time_reporting.modules.users.contracts import UserDTO

# A month far from any other test's data.
YEAR, MONTH = 2031, 4


@pytest.fixture(autouse=True)
def _isolated(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    patched = get_settings().model_copy(update={"attachment_dir": str(tmp_path)})
    monkeypatch.setattr("time_reporting.modules.purchases.service.get_settings", lambda: patched)


@pytest.fixture
async def accountant(make_user: UserFactory) -> UserDTO:
    return await make_user(roles=ACCOUNTANT)


async def _register(
    bus: Bus, actor: UserDTO, details: PurchaseDetails, *, name: str = "scan.pdf"
) -> PurchaseDocumentDTO:
    inbox = await bus.execute(
        AddPurchaseDocument(
            actor_id=None, file_name=name, content_type="application/pdf", content=b"%PDF-1"
        )
    )
    return await bus.execute(
        RegisterPurchaseDocument(document_id=inbox.id, actor_id=actor.id, details=details)
    )


def _receipt(day: date, method: PaymentMethod, vendor: str) -> PurchaseDetails:
    return PurchaseDetails(
        kind=PurchaseKind.RECEIPT,
        vendor=vendor,
        document_date=day,
        amount=Decimal("10.00"),
        currency="SEK",
        payment_method=method,
    )


async def _seed_month(bus: Bus, actor: UserDTO) -> None:
    card_invoice = await _register(
        bus,
        actor,
        PurchaseDetails(
            kind=PurchaseKind.CARD_INVOICE,
            vendor="Bank Card",
            document_no="K-1",
            document_date=date(YEAR, MONTH, 2),
            due_date=date(YEAR, MONTH, 25),
            payment_status=PaymentStatus.UNPAID,
            amount=Decimal("50.00"),
            currency="SEK",
        ),
    )
    linked = await _register(
        bus, actor, _receipt(date(YEAR, MONTH - 1, 20), PaymentMethod.CARD, "Cafe")
    )
    await _register(bus, actor, _receipt(date(YEAR, MONTH, 5), PaymentMethod.CARD, "Loose Shop"))
    await _register(bus, actor, _receipt(date(YEAR, MONTH, 9), PaymentMethod.CASH, "Kiosk"))
    await _register(
        bus,
        actor,
        PurchaseDetails(
            kind=PurchaseKind.INVOICE,
            vendor="Supplier AB",
            document_no="F-9",
            document_date=date(YEAR, MONTH, 10),
            due_date=date(YEAR, MONTH, 30),
            payment_status=PaymentStatus.UNPAID,
            amount=Decimal("200.00"),
            currency="SEK",
        ),
    )
    await bus.execute(
        LinkCardReceipts(
            card_invoice_id=card_invoice.id,
            actor_id=actor.id,
            links=(CardReceiptLink(receipt_id=linked.id, amount_base=Decimal("9.50")),),
        )
    )


async def test_status_counts_purchases_totals_and_warnings(bus: Bus, accountant: UserDTO) -> None:
    await _seed_month(bus, accountant)
    await bus.execute(
        AddPurchaseDocument(
            actor_id=None, file_name="waiting.pdf", content_type="application/pdf", content=b"%PDF"
        )
    )

    status = await bus.query(GetAccountantPackageStatus(year=YEAR, month=MONTH))

    # Card invoice + loose card receipt + cash receipt + supplier invoice; the linked receipt
    # (dated in March) belongs to the card invoice and isn't counted again.
    assert status.purchase_count == 4
    assert status.purchase_base_currency == "SEK"
    # 50 (card invoice) + 10 (cash) + 200 (supplier invoice); the loose card receipt has no
    # base amount yet.
    assert status.purchase_total_base == Decimal("260.00")
    assert status.purchase_total_provisional is True  # the unpaid amounts are estimates
    assert status.purchase_unconverted_count == 1
    assert status.unlinked_card_receipt_count == 1
    assert status.inbox_document_count >= 1


async def test_package_lists_purchases_and_groups_card_receipts(
    bus: Bus, accountant: UserDTO
) -> None:
    await _seed_month(bus, accountant)

    package = await bus.query(
        BuildAccountantPackage(year=YEAR, month=MONTH, actor_id=accountant.id)
    )
    try:
        with zipfile.ZipFile(package.path) as archive:
            purchase_files = sorted(n for n in archive.namelist() if n.startswith("purchases/"))
            assert purchase_files == [
                "purchases/003_2031-04-05_Loose-Shop_10.00-SEK.pdf",
                "purchases/004_2031-04-09_Kiosk_10.00-SEK.pdf",
                "purchases/005_2031-04-10_Supplier-AB_200.00-SEK.pdf",
                "purchases/card-2031-04-02/001_2031-04-02_Bank-Card_50.00-SEK.pdf",
                "purchases/card-2031-04-02/002_2031-03-20_Cafe_10.00-SEK.pdf",
            ]

            sheet = load_workbook(BytesIO(archive.read("summary.xlsx")))["Purchases"]
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            vendors = [row[3] for row in rows]
            assert vendors == ["Bank Card", "Cafe", "Loose Shop", "Kiosk", "Supplier AB"]
            cafe = rows[1]
            assert cafe[12] == "K-1"  # grouped under its card invoice
            assert cafe[8] == 9.5  # base amount typed from the card invoice
            assert rows[0][12] in (None, "")
            assert rows[4][9] == "Yes"  # the unpaid supplier invoice is provisional
            assert rows[4][10] == "due 2031-04-30"
            assert rows[2][8] is None  # the loose card receipt has no base amount

            assert archive.read("summary.pdf").startswith(b"%PDF")
    finally:
        package.path.unlink(missing_ok=True)
